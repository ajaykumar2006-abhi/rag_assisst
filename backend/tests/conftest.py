"""Shared pytest fixtures.

Every test runs against a throwaway data directory: an isolated SQLite file, an
isolated upload folder and an isolated vector index. No test touches the user's
real library, and no test needs a network connection or an API key.

The default service container is injected directly onto ``app.state`` instead of
going through the FastAPI lifespan handler, so the fixture controls exactly which
configuration the API sees.
"""

from __future__ import annotations

import os
import sys
import tempfile
import time
from pathlib import Path
from typing import Callable, Dict, List, Optional

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

# Point application defaults at a scratch directory *before* importing app code,
# so importing create_app() cannot create directories inside the user's project.
os.environ.setdefault("DATA_DIR", tempfile.mkdtemp(prefix="rag-test-default-"))

from fastapi.testclient import TestClient  # noqa: E402

from app.core.config import Settings  # noqa: E402
from app.main import create_app  # noqa: E402
from app.services.container import ServiceContainer  # noqa: E402

POLL_INTERVAL_SECONDS = 0.1
POLL_TIMEOUT_SECONDS = 60.0


# ---------------------------------------------------------------------------
# Core fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    """A Settings clone that writes only inside ``tmp_path``."""
    config = Settings()
    config.data_dir = tmp_path / "data"
    config.ensure_dirs()
    # Keep tests fast and deterministic.
    config.chunk_size = 400
    config.chunk_overlap = 60
    config.top_k = 5
    config.answer_provider = "extractive"
    config.embedding_provider = "hashing"
    return config


@pytest.fixture
def container(settings: Settings):
    """A real service container bound to the temporary settings."""
    instance = ServiceContainer(settings)
    try:
        yield instance
    finally:
        instance.shutdown()


@pytest.fixture
def client(container: ServiceContainer):
    """TestClient wired to the fixture container.

    The app is deliberately *not* entered as a context manager: that would run the
    lifespan handler, which would build (and leak) a second container from the
    ambient configuration instead of the isolated one the test owns.
    """
    app = create_app()
    app.state.container = container
    yield TestClient(app)
    app.state.container = container


# ---------------------------------------------------------------------------
# Document factories
# ---------------------------------------------------------------------------

HELIOS_TEXT = (
    "Project Helios was approved in March 2031 by the orbital review board.\n\n"
    "The mission budget was 418 million euros for the prime phase.\n\n"
    "Its primary instrument is the Cryo-Spectrometer Array, built by the Tycho Institute.\n\n"
    "The observatory operates at an altitude of 612 kilometres with an inclination of 28.5 degrees.\n"
)

LONG_TEXT = "\n\n".join(
    f"Section {index}: the deep field survey recorded {index * 7} calibration events "
    f"during orbital pass number {index}. The spectrometer logged "
    f"{index} interference flags and the recorder buffered {index * 3} hours of data."
    for index in range(1, 60)
)


def make_txt(name: str = "helios.txt", text: str = HELIOS_TEXT) -> tuple[str, bytes, str]:
    return name, text.encode("utf-8"), "text/plain"


def make_pdf(name: str = "helios.pdf", pages: Optional[List[str]] = None) -> tuple[str, bytes, str]:
    """Build a small multi-page PDF with real, extractable text."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    import io

    pages = pages or [HELIOS_TEXT, "Page two mentions the Fine Guidance Sensor at 0.18 arcseconds."]
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for page_text in pages:
        text_object = pdf.beginText(40, 780)
        text_object.setFont("Helvetica", 10)
        for line in page_text.split("\n"):
            text_object.textLine(line)
        pdf.drawText(text_object)
        pdf.showPage()
    pdf.save()
    return name, buffer.getvalue(), "application/pdf"


def make_docx(name: str = "helios.docx", paragraphs: Optional[List[str]] = None) -> tuple[str, bytes, str]:
    import docx
    import io

    paragraphs = paragraphs or [
        "DOCX ingestion fixture.",
        "The anomaly review board meets weekly and classifies each event.",
    ]
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Parameter"
    table.rows[0].cells[1].text = "Value"
    buffer = io.BytesIO()
    document.save(buffer)
    return (
        name,
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@pytest.fixture
def upload() -> Callable[..., Dict]:
    """Upload a file through the real HTTP endpoint."""
    def _upload(client: TestClient, file_tuple: tuple[str, bytes, str], field: str = "file"):
        name, data, content_type = file_tuple
        return client.post(
            "/api/v1/documents",
            files={field: (name, data, content_type)},
        )

    return _upload


@pytest.fixture
def wait_for_ready(client: TestClient) -> Callable[[str, float], Dict]:
    """Poll the status endpoint until processing settles."""
    def _wait(document_id: str, timeout: float = POLL_TIMEOUT_SECONDS) -> Dict:
        deadline = time.time() + timeout
        last: Dict = {}
        while time.time() < deadline:
            response = client.get(f"/api/v1/documents/{document_id}/status")
            assert response.status_code == 200, response.text
            last = response.json()
            if last["status"] in ("ready", "failed"):
                return last
            time.sleep(POLL_INTERVAL_SECONDS)
        raise AssertionError(f"Document {document_id} never settled; last status {last}")

    return _wait


@pytest.fixture
def ingest(client: TestClient, upload: Callable, wait_for_ready: Callable) -> Callable:
    """Upload a document and return (document_id, detail) once it is indexed."""
    def _ingest(file_tuple: tuple[str, bytes, str], expect_ready: bool = True):
        response = upload(client, file_tuple)
        assert response.status_code == 201, response.text
        document_id = response.json()["document"]["id"]
        status = wait_for_ready(document_id)
        if expect_ready:
            assert status["status"] == "ready", status
        detail = client.get(f"/api/v1/documents/{document_id}").json()
        return document_id, detail

    return _ingest
