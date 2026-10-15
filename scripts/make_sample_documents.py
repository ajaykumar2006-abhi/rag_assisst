"""Generate realistic sample documents for manual and browser testing.

Creates, from the synthetic Helios evaluation corpus in ``eval/datasets``:

* ``eval/fixtures/helios_mission_brief.pdf``      - three pages with a real text layer
* ``eval/fixtures/helios_instrument_specs.txt``   - plain text
* ``eval/fixtures/helios_ops_procedures.docx``    - Word document with a table

These are genuine PDF/TXT/DOCX files (not renamed text), so they exercise the
real extractors, and the PDF carries page breaks so citation page numbers can be
verified against the document the answer came from.

Usage
-----
    python scripts/make_sample_documents.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATASET = PROJECT_ROOT / "eval" / "datasets" / "synthetic_galaxy.json"
FIXTURE_DIR = PROJECT_ROOT / "eval" / "fixtures"

PAGE_WIDTH_LINES = 46


def wrap(text: str, width: int = 92) -> list[str]:
    """Wrap a paragraph to a fixed character width, preserving blank lines."""
    lines: list[str] = []
    for paragraph in text.split("\n"):
        if not paragraph.strip():
            lines.append("")
            continue
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}".strip()
            if len(candidate) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def paginate(lines: list[str]) -> list[list[str]]:
    """Split wrapped lines into pages of at most PAGE_WIDTH_LINES lines."""
    return [lines[index : index + PAGE_WIDTH_LINES] for index in range(0, len(lines), PAGE_WIDTH_LINES)]


def write_pdf(path: Path, documents: list[dict]) -> int:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pages: list[list[str]] = []
    for document in documents:
        pages.extend(paginate(wrap(document["text"])))

    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle("Helios Deep Field Survey - Mission Brief")
    for number, page_lines in enumerate(pages, start=1):
        pdf.setFont("Helvetica", 9)
        text = pdf.beginText(46, 800)
        for line in page_lines:
            text.textLine(line)
        pdf.drawText(text)
        pdf.setFont("Helvetica", 8)
        pdf.drawString(46, 30, f"Helios Deep Field Survey - page {number} of {len(pages)}")
        pdf.showPage()
    pdf.save()
    return len(pages)


def write_txt(path: Path, document: dict) -> None:
    path.write_text(document["text"] + "\n", encoding="utf-8")


def write_docx(path: Path, document: dict) -> None:
    import docx

    result = docx.Document()
    result.add_heading("Helios Deep Field Survey - Operations Procedures", level=1)
    for paragraph in document["text"].split("\n\n"):
        result.add_paragraph(paragraph.strip())

    table = result.add_table(rows=4, cols=2)
    rows = [
        ("Parameter", "Value"),
        ("Downlink passes per day", "4"),
        ("Nominal pass duration", "11 minutes"),
        ("Calibration interval", "30 days"),
    ]
    for index, (left, right) in enumerate(rows):
        table.rows[index].cells[0].text = left
        table.rows[index].cells[1].text = right

    result.save(str(path))


def main() -> int:
    if not DATASET.exists():
        print(f"Dataset not found: {DATASET}", file=sys.stderr)
        return 1

    dataset = json.loads(DATASET.read_text(encoding="utf-8"))
    by_id = {document["id"]: document for document in dataset["documents"]}

    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)

    pdf_path = FIXTURE_DIR / "helios_mission_brief.pdf"
    page_count = write_pdf(pdf_path, [by_id["overview"], by_id["instruments"]])

    txt_path = FIXTURE_DIR / "helios_instrument_specs.txt"
    write_txt(txt_path, by_id["instruments"])

    docx_path = FIXTURE_DIR / "helios_ops_procedures.docx"
    write_docx(docx_path, by_id["operations"])

    print(f"Wrote {pdf_path} ({pdf_path.stat().st_size} bytes, {page_count} pages)")
    print(f"Wrote {txt_path} ({txt_path.stat().st_size} bytes)")
    print(f"Wrote {docx_path} ({docx_path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
